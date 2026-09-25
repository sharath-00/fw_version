import json
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

data = json.load(open('latest_fw_data.json'))
# Filter range 1 to 500 and only fw_version .55
filtered_panels = []
for p in data:
    name = p.get('panel_name', '')
    if name.startswith('SALSS'):
        num_str = name[5:]
        if num_str.isdigit():
            if 1 <= int(num_str) <= 500 and '.55' in p.get('fw_version', ''):
                filtered_panels.append(p)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Panels_001_to_500'

# Headers
headers = ['Panel ID', 'Region', 'Zone', 'Ward', 'Current FW Version', 'New FW Version']
ws.append(headers)

# Populate data
for p in sorted(filtered_panels, key=lambda x: x.get('panel_name', '')):
    ws.append([
        p.get('panel_name', ''),
        p.get('region', ''),
        p.get('zone', ''),
        p.get('ward', ''),
        p.get('fw_version', ''),
        p.get('fw_version', '') # This will be the dropdown column
    ])

# Data validation (dropdown) for column E and F
dv = DataValidation(type='list', formula1='"SL530.54,SL530.55,SL530.47"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'F2:F{len(filtered_panels)+1}')
dv.add(f'E2:E{len(filtered_panels)+1}')

# Conditional Formatting
red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

ws.conditional_formatting.add(f'F2:F{len(filtered_panels)+1}', CellIsRule(operator='equal', formula=['"SL530.55"'], stopIfTrue=True, fill=red_fill))
ws.conditional_formatting.add(f'F2:F{len(filtered_panels)+1}', CellIsRule(operator='equal', formula=['"SL530.54"'], stopIfTrue=True, fill=green_fill))
ws.conditional_formatting.add(f'E2:E{len(filtered_panels)+1}', CellIsRule(operator='equal', formula=['"SL530.55"'], stopIfTrue=True, fill=red_fill))
ws.conditional_formatting.add(f'E2:E{len(filtered_panels)+1}', CellIsRule(operator='equal', formula=['"SL530.54"'], stopIfTrue=True, fill=green_fill))

# Auto-adjust column widths
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter # Get the column name
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column].width = adjusted_width

# Output into the workspace directory
wb.save('Panels_FW_001_to_500_Filtered.xlsx')
print('Excel generated: Panels_FW_001_to_500_Filtered.xlsx')
